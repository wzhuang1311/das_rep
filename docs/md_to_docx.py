r"""
md_to_docx.py —— 把复现报告 Markdown 转成 Word 文档

跑法：python md_to_docx.py
输入：das_rep\docs\复现报告.md
输出：das_rep\docs\复现报告.docx
"""
import os
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
REP = os.path.dirname(HERE)          # das_rep
MD = os.path.join(REP, "docs", "复现报告.md")
OUT = os.path.join(REP, "docs", "复现报告.docx")

CN_FONT = "微软雅黑"
EN_FONT = "Segoe UI"


def set_run_font(run, size=None, bold=None, cn=CN_FONT, en=EN_FONT):
    run.font.name = en
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold


def add_para(doc, text, size=10.5, bold=False, style=None, align=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    # 处理行内的 **粗体** 和 `代码`
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            set_run_font(r, size=size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1])
            set_run_font(r, size=size - 0.5, cn="Consolas", en="Consolas")
            r.font.color.rgb = RGBColor(0xB0, 0x30, 0x60)
        else:
            r = p.add_run(part)
            set_run_font(r, size=size, bold=bold)
    return p


def parse_table(lines, start):
    """解析 markdown 表格，返回 (表头, 数据行, 下一行索引)"""
    header = [c.strip() for c in lines[start].strip().strip("|").split("|")]
    rows = []
    i = start + 2                      # 跳过分隔行
    while i < len(lines) and lines[i].strip().startswith("|"):
        rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
        i += 1
    return header, rows, i


def add_table(doc, header, rows, avail_in=6.9):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    t.autofit = False
    # 按各列内容长度分配宽度，保证总宽不超出页面可用宽度
    ncol = len(header)
    lens = []
    for j in range(ncol):
        L = len(str(header[j]))
        for row in rows:
            if j < len(row):
                L = max(L, len(str(row[j]).replace("**", "")))
        lens.append(max(L, 4))
    total = sum(lens)
    widths = [max(0.55, avail_in * l / total) for l in lens]
    # 归一化，确保总和不超过可用宽度
    scale = avail_in / sum(widths)
    widths = [w * scale for w in widths]
    for j, w in enumerate(widths):
        for cell in t.columns[j].cells:
            cell.width = Inches(w)

    def fill(cell, txt, bold=False):
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        r = p.add_run(txt)
        set_run_font(r, size=8.5, bold=bold)

    for j, h in enumerate(header):
        fill(t.rows[0].cells[j], h, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for j in range(ncol):
            v = row[j] if j < len(row) else ""
            fill(cells[j], v.replace("**", ""))
    doc.add_paragraph()


def main():
    lines = open(MD, encoding="utf-8").read().split("\n")
    doc = Document()

    # 页面设置
    sec = doc.sections[0]
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)

    # 正文默认字体
    style = doc.styles["Normal"]
    style.font.name = EN_FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

    i = 0
    in_code = False
    code_buf = []
    while i < len(lines):
        line = lines[i]
        s = line.strip()

        # 代码块
        if s.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                r = p.add_run("\n".join(code_buf))
                set_run_font(r, size=8.5, cn="Consolas", en="Consolas")
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.space_after = Pt(6)
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # 空行
        if not s:
            i += 1
            continue

        # 分隔线
        if re.match(r"^-{3,}$", s):
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            lvl = len(m.group(1))
            txt = m.group(2).strip()
            if lvl == 1:
                h = doc.add_heading("", level=0)
            else:
                h = doc.add_heading("", level=min(lvl - 0, 4))
            r = h.add_run(txt)
            set_run_font(r, size={1: 20, 2: 15, 3: 13, 4: 11.5}.get(lvl, 11), bold=True)
            i += 1
            continue

        # 表格
        if s.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s\-:|]+\|$", lines[i+1].strip()):
            header, rows, ni = parse_table(lines, i)
            add_table(doc, header, rows)
            i = ni
            continue

        # 引用
        if s.startswith(">"):
            p = add_para(doc, s.lstrip("> ").strip(), size=10)
            p.paragraph_format.left_indent = Inches(0.3)
            for r in p.runs:
                r.font.italic = True
            i += 1
            continue

        # 列表项
        m = re.match(r"^(\d+)\.\s+(.*)$", s)
        if m:
            add_para(doc, "%s. %s" % (m.group(1), m.group(2)), size=10.5, style="List Number")
            i += 1
            continue
        if s.startswith("- "):
            add_para(doc, s[2:], size=10.5, style="List Bullet")
            i += 1
            continue

        # 普通段落
        add_para(doc, s)
        i += 1

    doc.save(OUT)
    print("已生成:", OUT)
    print("段落数:", len(doc.paragraphs), " 表格数:", len(doc.tables))


if __name__ == "__main__":
    main()
